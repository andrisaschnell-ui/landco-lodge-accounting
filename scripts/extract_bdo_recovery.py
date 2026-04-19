import json
import os
from datetime import date, datetime

import openpyxl


ROOT = r"C:\Users\Andrisa\Documents\Projects\landco"
BDO_DIR = os.path.join(ROOT, "Accounts 2026", "BDO Accounts 2026")


def parse_date(value):
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value).strip()
    if not text:
        return None
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            pass
    try:
        return datetime.fromisoformat(text).date().isoformat()
    except ValueError:
        return None


def num(value):
    if value in (None, ""):
        return 0
    try:
        return float(value)
    except Exception:
        try:
            return float(str(value).replace(",", "").strip())
        except Exception:
            return 0


def extract_petty_cash(file_path):
    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True, keep_links=False)
    try:
        ws = wb["Petty cash "]
        month = int(os.path.basename(file_path)[:2])
        rows = []
        started = False

        for row in ws.iter_rows(min_row=11, values_only=True):
            date_value = parse_date(row[0] if len(row) > 0 else None)
            reference = str(row[1]).strip() if len(row) > 1 and row[1] not in (None, "") else None
            supplier = str(row[2]).strip() if len(row) > 2 and row[2] not in (None, "") else None
            description = str(row[3]).strip() if len(row) > 3 and row[3] not in (None, "") else None
            allocation = str(row[4]).strip() if len(row) > 4 and row[4] not in (None, "") else None
            credit = num(row[5] if len(row) > 5 else None)
            debit = num(row[6] if len(row) > 6 else None)
            balance = num(row[7] if len(row) > 7 else None)
            vat_amount = num(row[8] if len(row) > 8 else None)
            net_amount = num(row[9] if len(row) > 9 else None)

            if not any([date_value, supplier, description, credit, debit]):
                if started:
                    break
                continue

            if not date_value:
                continue

            started = True
            combined_description = " | ".join(part for part in [supplier, description, allocation] if part) or "Petty cash item"

            rows.append(
                {
                    "date": date_value,
                    "description": combined_description,
                    "credit": credit,
                    "debit": debit,
                    "balance": balance,
                    "month": month,
                    "year": 2026,
                    "reference": reference,
                    "supplier": supplier,
                    "allocation": allocation,
                    "vat_amount": vat_amount,
                    "net_amount": net_amount,
                    "source_file": os.path.basename(file_path),
                }
            )

        return rows
    finally:
        wb.close()


def extract_creditors(file_path):
    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True, keep_links=False)
    try:
        if "Creditors" not in wb.sheetnames:
            return []
        ws = wb["Creditors"]
        rows = []
        started = False

        # Data starts around row 15 based on inspection
        for row in ws.iter_rows(min_row=13, values_only=True):
            if row[0] == "No" or row[0] is None:
                continue
            
            invoice_date = parse_date(row[1])
            invoice_num = str(row[2]).strip() if row[2] else None
            supplier = str(row[3]).strip() if row[3] else None
            description = str(row[4]).strip() if row[4] else None
            allocation = str(row[5]).strip() if row[5] else None
            amount_excl = num(row[6])
            vat_amount = num(row[7])
            total_amount = num(row[8])

            if not supplier and not amount_excl:
                if started: break
                continue
            
            started = True
            rows.append({
                "invoice_date": invoice_date,
                "invoice_number": invoice_num,
                "supplier": supplier,
                "description": description,
                "allocation": allocation,
                "amount_excl": amount_excl,
                "vat_amount": vat_amount,
                "total_amount": total_amount,
                "source_file": os.path.basename(file_path)
            })

        return rows
    finally:
        wb.close()


def extract_bank_tx(file_path):
    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True, keep_links=False)
    try:
        results = {"mtn": [], "usd": []}
        
        # Process MZN
        if "BIM Bank Control Mtn" in wb.sheetnames:
            ws = wb["BIM Bank Control Mtn"]
            started = False
            for row in ws.iter_rows(min_row=12, values_only=True):
                if not row[1] and not row[4]: # Date and Recipient check
                    if started: break
                    continue
                started = True
                results["mtn"].append({
                    "date": parse_date(row[1]),
                    "reference": str(row[2]).strip() if row[2] else None,
                    "description": str(row[3]).strip() if row[3] else None,
                    "recipient": str(row[4]).strip() if row[4] else None,
                    "allocation": str(row[5]).strip() if row[5] else None,
                    "amount_excl": num(row[6]),
                    "vat_amount": num(row[7]),
                    "amount_incl": num(row[8]),
                    "source_file": os.path.basename(file_path)
                })
        
        # Process USD
        if "BIM Bank Control USD" in wb.sheetnames:
            ws = wb["BIM Bank Control USD"]
            started = False
            for row in ws.iter_rows(min_row=12, values_only=True):
                if not row[1] and not row[4]:
                    if started: break
                    continue
                started = True
                results["usd"].append({
                    "date": parse_date(row[1]),
                    "reference": str(row[2]).strip() if row[2] else None,
                    "description": str(row[3]).strip() if row[3] else None,
                    "recipient": str(row[4]).strip() if row[4] else None,
                    "allocation": str(row[5]).strip() if row[5] else None,
                    "amount_excl": num(row[6]),
                    "vat_amount": num(row[7]),
                    "amount_incl": num(row[8]),
                    "source_file": os.path.basename(file_path)
                })
        
        return results
    finally:
        wb.close()


def extract_equity(file_path):
    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True, keep_links=False)
    try:
        if "Shareholders investment" not in wb.sheetnames:
            return []
        ws = wb["Shareholders investment"]
        rows = []
        started = False
        for row in ws.iter_rows(min_row=12, values_only=True):
            if not row[1] and not row[2]:
                if started: break
                continue
            if row[1] == "DATE": continue
            
            started = True
            rows.append({
                "date": parse_date(row[1]),
                "shareholder": str(row[2]).strip() if row[2] else None,
                "description": str(row[3]).strip() if row[3] else None,
                "credit": num(row[4]),
                "debit": num(row[5]),
                "balance": num(row[6]),
                "source_file": os.path.basename(file_path)
            })
        return rows
    finally:
        wb.close()


def main():
    files = sorted(
        os.path.join(BDO_DIR, name)
        for name in os.listdir(BDO_DIR)
        if name.lower().endswith(".xlsx") and not name.startswith("~$")
    )
    payload = {"petty_cash": [], "creditors": [], "bank_outflows": {"mtn": [], "usd": []}, "equity": []}
    for file_path in files:
        payload["petty_cash"].extend(extract_petty_cash(file_path))
        payload["creditors"].extend(extract_creditors(file_path))
        
        btx = extract_bank_tx(file_path)
        payload["bank_outflows"]["mtn"].extend(btx["mtn"])
        payload["bank_outflows"]["usd"].extend(btx["usd"])
        
        payload["equity"].extend(extract_equity(file_path))
        
    print(json.dumps(payload))


if __name__ == "__main__":
    main()
